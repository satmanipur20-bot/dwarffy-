"""Render a :class:`ClientData` into the branded one-page Requirement List.

The branded template (``templates/SAT_Requirement_Template.xlsx``) carries the
embedded logo and the fixed letterhead (rows 1-13).  We load it, fill in the
assessee block / due-date bar / preamble, then build the data region (Section A
banks, Section B investments, closing note) from scratch with explicit styling
that matches the firm's house style.
"""

from __future__ import annotations

import os
from typing import List

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .forms import FormDecision, determine_form_and_due_date
from .model import BankAccount, ClientData, Investment

# ---- brand palette -------------------------------------------------------
CHARCOAL = "FF404042"
ORANGE = "FFEB8025"
LIGHT_ORANGE = "FFFBE3CC"
LABEL = "FFECECEC"
ALT_ROW = "FFF4F4F4"
WHITE = "FFFFFFFF"

FONT = "Calibri"

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..",
                             "templates", "SAT_Requirement_Template.xlsx")

# rows fixed by the template scaffold
ROW_NAME, ROW_STATUS, ROW_ADDRESS = 7, 8, 9
ROW_DUE_DATE = 10
ROW_PREAMBLE = 11
ROW_TABLE_HEADER = 13
FIRST_BODY_ROW = 14

CLOSING_NOTE = ("Note: The above list is indicative. Kindly intimate any new bank account, "
                "investment, income, asset, or deduction not appearing above so that the same "
                "may be appropriately considered.")


def _side(style):
    return Side(style=style, color=CHARCOAL)


def _body_border(col: str, last: bool) -> Border:
    """Outer edges medium, inner thin; bottom of the final data row medium."""
    left = _side("medium") if col == "A" else _side("thin")
    right = _side("medium") if col == "D" else _side("thin")
    bottom = _side("medium") if last else _side("thin")
    return Border(left=left, right=right, top=_side("thin"), bottom=bottom)


def _investment_label(inv: Investment) -> str:
    """Format column B for a structured investment (or pass ``label`` through)."""
    if inv.label:
        return inv.label
    k, e = inv.kind, inv.entity
    branch = f" ({inv.branch})" if inv.branch else ""
    if k == "fixed_deposit":
        return f"Fixed Deposit – {e}{branch}"
    if k == "recurring_deposit":
        return f"Recurring Deposit – {e}{branch}"
    if k == "mutual_fund":
        via = f" (via {inv.rta})" if inv.rta else ""
        return f"Mutual Fund – {e} AMC{via}"
    if k == "kvp":
        return f"Kisan Vikas Patra (KVP){(' – ' + e) if e else ''}"
    if k == "nsc":
        return f"National Savings Certificate (NSC){(' – ' + e) if e else ''}"
    if k == "bond":
        return f"Bonds – {e}" if e else "Bonds"
    if k == "share":
        return "Shares & Securities – Demat Holding & Transaction Statement"
    if k == "demat":
        return "Shares & Securities – Demat Holding & Transaction Statement"
    if k == "lic":
        return f"LIC – {e}" if e else "LIC Policy (Premium Receipt / Statement)"
    if k == "gold":
        return "Gold & Jewellery (if made any investment)"
    return e or "Investment"


def _write_body_cell(ws, row, col_letter, value, *, fill, align_h,
                     bold=False, size=10.5, color="FF000000", text_fmt=False,
                     last=False):
    cell = ws[f"{col_letter}{row}"]
    if text_fmt:
        cell.number_format = "@"          # set BEFORE assigning, preserves zeros
    cell.value = value
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=True)
    cell.border = _body_border(col_letter, last)
    return cell


def _section_row(ws, row, title):
    for col in "ABCD":
        _write_body_cell(ws, row, col, title if col == "A" else None,
                         fill=ORANGE, align_h="left", bold=True, size=11,
                         color=CHARCOAL)


def _data_row(ws, row, serial, particulars, ref, *, last=False):
    fill = WHITE if serial % 2 == 1 else ALT_ROW   # odd serial → white
    _write_body_cell(ws, row, "A", serial, fill=fill, align_h="center", last=last)
    _write_body_cell(ws, row, "B", particulars, fill=fill, align_h="left", last=last)
    _write_body_cell(ws, row, "C", ref or "—", fill=fill, align_h="center",
                     text_fmt=True, last=last)
    # Column D (Remarks / Action Required) is left BLANK in every data row.
    _write_body_cell(ws, row, "D", None, fill=fill, align_h="left", last=last)


def _madam_or_sir(gender: str) -> str:
    return "Madam" if (gender or "").strip().lower().startswith("f") else "Sir"


def build_requirement_xlsx(client: ClientData, output_path: str,
                           template_path: str = TEMPLATE_PATH) -> FormDecision:
    """Build the workbook and return the form/due-date decision (for the memo)."""
    p = client.particulars
    decision = determine_form_and_due_date(client.profile, p.assessment_year)

    wb = load_workbook(template_path)
    ws = wb["Requirement"]

    # ---- assessee block -------------------------------------------------
    ws[f"B{ROW_NAME}"] = p.name
    ws[f"D{ROW_NAME}"] = p.pan
    ws[f"B{ROW_STATUS}"] = p.status
    ws[f"D{ROW_STATUS}"] = p.assessment_year
    ws[f"B{ROW_ADDRESS}"] = p.address
    ws[f"D{ROW_ADDRESS}"] = p.previous_year

    # ---- due-date bar ---------------------------------------------------
    ws[f"A{ROW_DUE_DATE}"] = (
        f"DUE DATE FOR FILING OF RETURN (A.Y. {p.assessment_year}):  {decision.due_date}"
        "   —   Kindly furnish all documents well in advance to enable timely filing."
    )

    # ---- preamble -------------------------------------------------------
    ay = p.assessment_year
    fy = p.previous_year.split(" ")[0] if p.previous_year else ""
    ws[f"A{ROW_PREAMBLE}"] = (
        f"Dear {_madam_or_sir(p.gender)},\n"
        f"With reference to the preparation and filing of your Income-Tax Return for the "
        f"Assessment Year {ay}, we request you to kindly furnish the following documents and "
        f"information at the earliest. The list has been compiled on the basis of your Annual "
        f"Information Statement (AIS) for FY {fy} and the financial statements of the preceding year."
    )

    # ---- body: Section A, Section B, closing note -----------------------
    row = FIRST_BODY_ROW
    serial = 0
    total_data_rows = len(client.bank_accounts) + len(client.investments)

    def is_last(s):
        return s == total_data_rows

    _section_row(ws, row, "A.   Bank Account Statements")
    ws.row_dimensions[row].height = 21.0
    row += 1
    for acc in client.bank_accounts:
        serial += 1
        label = f"{acc.bank} – {acc.account_type} Account" if acc.account_type else acc.bank
        if acc.branch:
            label += f" ({acc.branch} Branch)"
        _data_row(ws, row, serial, label, acc.account_no, last=is_last(serial))
        ws.row_dimensions[row].height = 21.0
        row += 1

    _section_row(ws, row, "B.   Investments")
    ws.row_dimensions[row].height = 21.0
    row += 1
    for inv in client.investments:
        serial += 1
        _data_row(ws, row, serial, _investment_label(inv), inv.ref,
                  last=is_last(serial))
        ws.row_dimensions[row].height = 21.0
        row += 1

    # ---- closing note ---------------------------------------------------
    note_row = row
    ws.merge_cells(f"A{note_row}:D{note_row}")
    note = ws[f"A{note_row}"]
    note.value = CLOSING_NOTE
    note.font = Font(name=FONT, size=9.5, italic=True, color=CHARCOAL)
    note.fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
    note.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[note_row].height = 30.0

    # ---- one landscape page --------------------------------------------
    ws.print_area = f"A1:D{note_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.scale = None
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False

    wb.save(output_path)
    return decision
